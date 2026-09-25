import pandas as pd
import openpyxl


def load_excel(file_path):
    """
    Load all sheets from an Excel workbook.
    """

    workbook = openpyxl.load_workbook(
        file_path,
        data_only=False
    )

    sheets = {}

    for sheet_name in workbook.sheetnames:

        worksheet = workbook[sheet_name]

        rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        sheets[sheet_name] = rows

    return sheets


def create_dataframe_context(file):

    excel_file = pd.ExcelFile(file)
    
    context = {}

    for sheet_name in excel_file.sheet_names:

        df = pd.read_excel(
            file_path,
            sheet_name=sheet_name
        )

        context[sheet_name] = df

    return context